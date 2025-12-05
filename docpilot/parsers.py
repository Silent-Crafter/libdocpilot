import pandas as pd
import openpyxl

from llama_index.core.readers.base import BaseReader
from llama_index.core.schema import Document
from pathlib import Path
from docpilot.preprocessors.pdf import PDFPreprocessor

from typing import Any, Dict, List, Optional

from docpilot.notlogging.notlogger import NotALogger

import pprint

logger = NotALogger(__name__)
logger.enabled = False

class CustomXLSXReader(BaseReader):
    def __init__(
            self,
            *args: Any,
            concat_rows: bool = True,
            col_joiner: str = ",",
            row_joiner: str = "\n",
            pandas_config: dict = None,
            **kwargs: Any
    ) -> None:
        """Init params."""
        super().__init__(*args, **kwargs)
        if pandas_config is None:
            pandas_config = {}
        self._concat_rows = concat_rows
        self._col_joiner = col_joiner
        self._row_joiner = row_joiner
        self._pandas_config = pandas_config

    def load_data(
            self,
            file: Path,
            extra_info: Optional[Dict] = None,
    ) -> List[Document]:
        """Parse file."""

        wb = openpyxl.load_workbook(file)

        docs: list[Document] = []
        for ws in wb.sheetnames:
            df = pd.read_excel(file, sheet_name=ws, **self._pandas_config)

            text_list = [" ".join(df.columns.astype(str))]  # Concat headers
            text_list += (
                df.astype(str)
                .apply(lambda row: self._col_joiner.join(row.values), axis=1)
                .tolist()
            )

            metadata = {"filename": file.name, "extension": file.suffix}
            if extra_info:
                metadata.update(extra_info)

            if self._concat_rows:
                docs.append(Document(
                    text=self._row_joiner.join(text_list),
                    metadata=metadata,
                ))
            else:
                docs.extend([
                    Document(
                        text=text,
                        metadata=metadata,
                    )
                    for text in text_list
                ])

        return docs


class CustomPDFReader(BaseReader):
    """PDF parser."""

    def __init__(self, return_full_document: Optional[bool] = False) -> None:
        """
        Initialize PDFReader.
        """
        self.return_full_document = return_full_document
        self.preprocessor = None

    def load_data(
            self,
            file: Path,
            extra_info: Optional[Dict] = None,
            use_artifact_pdf: bool = True,
    ) -> List[Document]:
        """Parse file."""
        if not isinstance(file, Path):
            file = Path(file)

        self.preprocessor=PDFPreprocessor(str(file))
        raw_elements=self.preprocessor.get_elements()
        self.preprocessor.close()

        documents=[]

        logger.info(f"Processing {file.name}")

        text_buffer=""
        buffer_page_start=1
        current_page=1

        base_metadata={
            "file_name":file.name,
            "create_date": file.stat().st_ctime_ns
        }

        def flush_buffer():
            nonlocal text_buffer, buffer_page_start

            if text_buffer.strip():
                doc=Document(
                    text=text_buffer.strip(),
                    metadata={
                        **base_metadata,
                        "type":"text"
                    }
                )
                documents.append(doc)
            text_buffer=""
        
        for item in raw_elements:
            if item["page"]>current_page:
                current_page=item["page"]
                text_buffer+=f"\n\n[PAGE {current_page}]\n\n"
            
            if item["type"]=="text":
                text_buffer+=item["content"] +"\n\n"
            
            elif item["type"]=="image":
                text_buffer+=f"\n[IMAGE REF: {item['content']}]\n"
                documents.append(Document(
                    text=item['content'],
                    metadata={
                        **base_metadata,
                        "page_number": item["page"],
                        "type":"image",
                        "image_path":item["image_path"],
                        "bbox": item["bbox"]
                    }
                ))
            
            elif item["type"]=="table":
                text_buffer+=f"\n[TABLE REF: {item['content']}]\n"
                documents.append(Document(
                    text=item['content'],
                    metadata={
                        **base_metadata,
                        "page_number": item["page"],
                        "type":"table",
                        "table_html":item["table_html"],
                        "bbox": item["bbox"]
                    }
                ))

        flush_buffer()

        return documents
    
if __name__=="__main__":
    pdfReader=CustomPDFReader()
    docs=pdfReader.load_data("test_data/Attention.pdf")
    pprint.pprint(docs[-1].text)