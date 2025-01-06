import sys
from pathlib import Path
from typing import Any, Dict, List, Optional
import pandas as pd
import openpyxl
from llama_index.core.readers.base import BaseReader
from llama_index.core.schema import Document
from llama_index.core.storage.index_store import SimpleIndexStore


class CustomXLSXReader(BaseReader):
    def __init__(
            self,
            *args: Any,
            concat_rows: bool = True,
            col_joiner: str = ",",
            row_joiner: str = "\n",
            pandas_config: dict = {},
            **kwargs: Any
    ) -> None:
        """Init params."""
        super().__init__(*args, **kwargs)
        self._concat_rows = concat_rows
        self._col_joiner = col_joiner
        self._row_joiner = row_joiner
        self._pandas_config = pandas_config

    def load_data(
            self,
            file: Path,
            extra_info: Optional[Dict] = None,
    ) -> List[Document]:
        """Parse file."""

        wb = openpyxl.load_workbook(file)

        documents: list[Document] = []
        for ws in wb.sheetnames:
            df = pd.read_excel(file, sheet_name=ws, **self._pandas_config)

            sys.stdout.flush()
            text_list = [" ".join(df.columns.astype(str))]  # Concat headers
            text_list += (
                df.astype(str)
                .apply(lambda row: self._col_joiner.join(row.values), axis=1)
                .tolist()
            )

            metadata = {"filename": file.name, "extension": file.suffix}
            if extra_info:
                metadata.update(extra_info)

            if self._concat_rows:
                documents.append(Document(
                    text=self._row_joiner.join(text_list),
                    metadata=metadata,
                ))
            else:
                documents.extend([
                    Document(
                        text=text,
                        metadata=metadata,
                    )
                    for text in text_list
                ])


        return documents


if __name__ == '__main__':
    from llama_index.core import SimpleDirectoryReader

    documents = SimpleDirectoryReader(
        "data",
        file_extractor={".xlsx": CustomXLSXReader()},
    ).load_data()

    # documents = CustomXLSXReader().load_data("data/machines.xlsx")

    print(documents[0].text)
