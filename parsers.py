import pandas as pd
import openpyxl

from llama_index.core.readers.base import BaseReader
from llama_index.core.schema import Document
from pathlib import Path
from preprocessors.pdf import PDFPreprocessor

from typing import Any, Dict, List, Optional

from notlogging.notlogger import NotALogger

logger = NotALogger(__name__)
logger.enable = False

class CustomXLSXReader(BaseReader):
    def __init__(
            self,
            *args: Any,
            concat_rows: bool = True,
            col_joiner: str = ",",
            row_joiner: str = "\n",
            pandas_config: dict = None,
            **kwargs: Any
    ) -> None:
        """Init params."""
        super().__init__(*args, **kwargs)
        if pandas_config is None:
            pandas_config = {}
        self._concat_rows = concat_rows
        self._col_joiner = col_joiner
        self._row_joiner = row_joiner
        self._pandas_config = pandas_config

    def load_data(
            self,
            file: Path,
            extra_info: Optional[Dict] = None,
    ) -> List[Document]:
        """Parse file."""

        wb = openpyxl.load_workbook(file)

        docs: list[Document] = []
        for ws in wb.sheetnames:
            df = pd.read_excel(file, sheet_name=ws, **self._pandas_config)

            text_list = [" ".join(df.columns.astype(str))]  # Concat headers
            text_list += (
                df.astype(str)
                .apply(lambda row: self._col_joiner.join(row.values), axis=1)
                .tolist()
            )

            metadata = {"filename": file.name, "extension": file.suffix}
            if extra_info:
                metadata.update(extra_info)

            if self._concat_rows:
                docs.append(Document(
                    text=self._row_joiner.join(text_list),
                    metadata=metadata,
                ))
            else:
                docs.extend([
                    Document(
                        text=text,
                        metadata=metadata,
                    )
                    for text in text_list
                ])

        return docs


class CustomPDFReader(BaseReader):
    """PDF parser."""

    def __init__(self, return_full_document: Optional[bool] = False) -> None:
        """
        Initialize PDFReader.
        """
        self.return_full_document = return_full_document

    def load_data(
            self,
            file: Path,
            extra_info: Optional[Dict] = None,
    ) -> List[Document]:
        """Parse file."""
        if not isinstance(file, Path):
            file = Path(file)

        metadata = {"file_name": file.name, "create_date": file.stat().st_ctime_ns}

        docs = []

        logger.log(f"Processing {file.name}", "DEBUG")
        preprocessor = PDFPreprocessor(file)
        pages = preprocessor.forward()

        text = ""
        for page in pages:
            page = list(filter(lambda x: x.strip(), page))
            text += "\n".join(page)
            text += "\n\n"

        # Join text extracted from each page
        docs.append(Document(text=text, metadata=metadata))

        return docs


if __name__ == '__main__':
    documents = CustomPDFReader().load_data(Path("data/Control Plan - 20. winding CP rev-28.pdf"))

    print(len(documents))
    print(documents[0].text)
